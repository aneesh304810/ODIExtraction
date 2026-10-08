#!/usr/bin/env python3
"""ODI 360 extractor - walks the PBDW Bitbucket clones and emits every data point
needed for an ODI 360 view (scenario -> load plan -> Autosys -> tables -> DB code).

Repos understood (clone them side by side, then point --root at the parent):
    oracle-data-integrator/   Master Repository/  Work Repository/
                              Scenarios and Load Plans/  Releases/
    prod-db-scripts/          Database Scripts/<SCHEMA>/{Functions,Index,Procedure,
                              Sequence,Tables,Views}
    Autosys/                  *.jil (insert_job ... command: startscen.sh ...)
    pb-dwh-scripts/, pbinboundprocess/, ... any other repo -> scanned as scripts

Usage:
    python odi360_extract.py --root D:\\bitbucket\\PBDW --out D:\\odi360_out
    python odi360_extract.py --odi-repo ..\\oracle-data-integrator ^
        --db-scripts ..\\prod-db-scripts --autosys ..\\Autosys ^
        --scripts ..\\pb-dwh-scripts --env PROD --out out

Outputs (under --out):
    datapoints/*.csv            one CSV per data-point family (raw facts)
    ODI360_datapoints.xlsx      same, one sheet per CSV (needs openpyxl)
    cp360/ODI_360.csv|.xlsx     CP360 ingest sheet (contract in prompts/)
    claude_ingest/              markdown knowledge + graph JSON for Claude Enterprise

Stdlib only (openpyxl optional). Parsing is heuristic and never fails the run:
unparseable files are listed in datapoints/parse_errors.csv.
"""
import argparse
import csv
import datetime
import io
import json
import os
import re
import subprocess
import sys
import zipfile
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict

VERSION = "1.0.0"

# ---------------------------------------------------------------- configuration
# Heuristic rules. Edit freely - first match wins. Tested against object names.
DOMAIN_RULES = [
    (r"ADDEPAR", "Addepar"),
    (r"BENCHMARK", "Benchmark"),
    (r"PERFORM", "Performance"),
    (r"HOLDING", "Holdings"),
    (r"POSITION|\bPOS_", "Positions"),
    (r"TRANS|\bTXN", "Transactions"),
    (r"ACCT|ACCOUNT", "Accounts"),
    (r"SETL|SETTLE", "Settlement"),
    (r"BILL", "Billpay"),
    (r"PRICE|PRICING", "Pricing"),
    (r"SECURIT|\bSEC_", "Securities"),
    (r"CLIENT|CUST", "Client"),
    (r"ASSET", "Assets"),
    (r"\bAPI|_API_", "API Ingestion"),
    (r"PURGE|ARCHIV|CLEAN|LOG_JOB", "Housekeeping"),
    (r"CHMOD|FILE|FTP", "File Handling"),
    (r"TEST", "Test / Utility"),
    (r"ADHOC", "Adhoc"),
    (r"INTRADAY", "Intraday"),
]
LAYER_RULES = [  # applied to target tables first, then the object name
    (r"(^|_)STG\d?_|^STG|STAGE", "Staging"),
    (r"^FACT_|_FACT_|^F_", "Warehouse Fact"),
    (r"^DIM_|_DIM_|^D_", "Warehouse Dimension"),
    (r"^BD_|MAPPING|^MAP_|^XREF", "Reference / Mapping"),
    (r"^RPT_|^MART|^AGG_|SMRY|SUMMARY", "Reporting / Mart"),
    (r"^API|_API_", "API Ingestion"),
    (r"^LOG|AUDIT|PURGE", "Operational"),
]
ODI_WORK_TABLE = re.compile(r"^(C\$|I\$|E\$|J\$|JV\$|T\$|SNP_|ODI_TMP|TMP_)", re.I)
SQL_RESERVED = {
    "DUAL", "SELECT", "TABLE", "WHERE", "SET", "VALUES", "THE", "LATERAL", "ONLY",
    "SYSDATE", "SYSTIMESTAMP", "DATE", "TIMESTAMP", "SESSION", "INTO", "AS", "ON",
    "NULL", "WITH", "EACH", "ROW", "TO", "BY", "OF", "OR", "AND", "IN", "IS",
}
JIRA_RE = re.compile(r"\b([A-Z][A-Z0-9]{1,9}-\d{1,6})\b")
TEXT_EXT = {".sql", ".pks", ".pkb", ".pck", ".prc", ".fnc", ".vw", ".trg", ".tab",
            ".seq", ".idx", ".sh", ".ksh", ".bash", ".bat", ".cmd", ".ps1", ".py",
            ".jil", ".txt", ".cfg", ".conf", ".properties", ".groovy", ".java",
            ".xml", ".json", ".yaml", ".yml", ".ctl", ".par", ""}
SKIP_DIRS = {".git", "node_modules", "__pycache__", ".idea", ".vscode", "target", "bin"}
MAX_CELL = 32000


# ---------------------------------------------------------------- small helpers
def g(o, *names, default=""):
    """First non-empty field among candidate ODI field names (versions differ)."""
    for n in names:
        v = o.get(n)
        if v not in (None, ""):
            return v.strip() if isinstance(v, str) else v
    return default


def norm(name):
    return (name or "").replace('"', "").strip().upper()


def short(name):
    """SCHEMA.TABLE@LINK -> TABLE"""
    n = norm(name).split("@")[0]
    return n.rsplit(".", 1)[-1]


def join(items, sep="; ", limit=MAX_CELL):
    s = sep.join(sorted({str(i) for i in items if i not in (None, "")}))
    return s if len(s) <= limit else s[: limit - 20] + " ...[truncated]"


def first_match(rules, *texts):
    for t in texts:
        if not t:
            continue
        for pat, label in rules:
            if re.search(pat, t, re.I):
                return label
    return ""


def read_text(path, max_bytes=50_000_000):
    try:
        with open(path, "rb") as f:
            raw = f.read(max_bytes)
    except OSError:
        return ""
    if b"\x00" in raw[:4096] and not raw.startswith(b"\xff\xfe") and not raw.startswith(b"\xfe\xff"):
        return ""  # binary
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return raw.decode("latin-1", "replace")


def walk(root, exts=None):
    for dp, dns, fns in os.walk(root):
        dns[:] = [d for d in dns if d not in SKIP_DIRS]
        for fn in sorted(fns):
            if exts is None or os.path.splitext(fn)[1].lower() in exts:
                yield os.path.join(dp, fn)


def rel(path, root):
    try:
        return os.path.relpath(path, root).replace("\\", "/")
    except ValueError:
        return path.replace("\\", "/")


# ---------------------------------------------------------------- git metadata
class GitIndex:
    """One `git log --name-only` pass per repo -> last commit + Jira keys per file."""

    def __init__(self, enabled=True):
        self.enabled = enabled
        self.by_path = {}
        self.done = set()

    def index(self, root):
        if not self.enabled or not root:
            return
        try:
            top = subprocess.run(["git", "-C", root, "rev-parse", "--show-toplevel"],
                                 capture_output=True, text=True, timeout=60).stdout.strip()
        except (OSError, subprocess.SubprocessError):
            return
        if not top or top in self.done:
            return
        self.done.add(top)
        try:
            out = subprocess.run(
                ["git", "-C", top, "log", "--name-only", "--no-renames", "--date=short",
                 "--format=%x1e%H%x1f%an%x1f%ad%x1f%s"],
                capture_output=True, text=True, encoding="utf-8", errors="replace",
                timeout=1800).stdout
        except (OSError, subprocess.SubprocessError):
            return
        for block in out.split("\x1e"):
            if not block.strip():
                continue
            head, _, files = block.partition("\n")
            parts = head.split("\x1f")
            if len(parts) < 4:
                continue
            sha, author, date, subject = parts[:4]
            jira = set(JIRA_RE.findall(subject))
            for f in files.splitlines():
                f = f.strip()
                if not f:
                    continue
                key = os.path.normcase(os.path.normpath(os.path.join(top, f)))
                rec = self.by_path.get(key)
                if rec is None:  # log is newest-first -> first sighting = last commit
                    rec = self.by_path[key] = {
                        "last_commit": sha[:10], "last_author": author, "last_date": date,
                        "last_subject": subject, "first_date": date, "n_commits": 0,
                        "jira": set(), "subjects": []}
                rec["n_commits"] += 1
                rec["first_date"] = date
                rec["jira"] |= jira
                if len(rec["subjects"]) < 5 and subject not in rec["subjects"]:
                    rec["subjects"].append(subject)

    def get(self, path):
        return self.by_path.get(os.path.normcase(os.path.normpath(os.path.abspath(path))), {})


# ---------------------------------------------------------------- SQL analysis
_COMMENT = re.compile(r"--[^\n]*|/\*.*?\*/", re.S)
_STRING = re.compile(r"'(?:[^']|'')*'")
_IDENT = r'(?:"[^"]+"|[A-Za-z_][\w$#]*)(?:\s*\.\s*(?:"[^"]+"|[A-Za-z_][\w$#]*))?(?:@[\w.$#]+)?'
_WRITES = [
    ("INSERT", re.compile(r"\binsert\s+(?:all\s+|first\s+)?into\s+(" + _IDENT + ")", re.I)),
    ("INSERT", re.compile(r"\binto\s+(" + _IDENT + r")\s*(?:\(|values|select)", re.I)),
    ("MERGE", re.compile(r"\bmerge\s+into\s+(" + _IDENT + ")", re.I)),
    ("UPDATE", re.compile(r"\bupdate\s+(" + _IDENT + r")\s+(?:\w+\s+)?set\b", re.I)),
    ("DELETE", re.compile(r"\bdelete\s+(?:from\s+)?(" + _IDENT + ")", re.I)),
    ("TRUNCATE", re.compile(r"\btruncate\s+table\s+(" + _IDENT + ")", re.I)),
    ("CREATE", re.compile(r"\bcreate\s+(?:global\s+temporary\s+)?table\s+(" + _IDENT + ")", re.I)),
    ("DROP", re.compile(r"\bdrop\s+table\s+(" + _IDENT + ")", re.I)),
]
_USING = re.compile(r"\busing\s+(" + _IDENT + ")", re.I)
_FROM_JOIN = re.compile(r"\b(from|join)\b", re.I)
_FROM_STOP = re.compile(
    r"\b(where|group|order|connect|start|union|minus|intersect|having|join|left|right|"
    r"inner|full|cross|natural|fetch|for|on|when|then|model|pivot|unpivot|partition|"
    r"sample|log|return|returning|into|set|values|select|as\s+of)\b", re.I)
_NOISE_FROM = re.compile(r"\b(extract|trim|substring)\s*\([^()]*?\bfrom\b", re.I)
_CALL = re.compile(r"\b(" + _IDENT + r")\s*\(")
_EXEC = re.compile(r"\b(?:exec(?:ute)?|call)\s+(" + _IDENT + ")", re.I)
_ODI_VAR = re.compile(r"#((?:GLOBAL|[A-Z][\w]*)\.[A-Z_][\w]*|[A-Z_][\w]{2,})", re.I)
_ODIREF = re.compile(r"odiRef\.(\w+)\(([^)]*)\)")
_JDBC = re.compile(r"jdbc:[\w:]+[@/][^\s'\"<>]+", re.I)
_HOST = re.compile(r"\b(?:[a-z0-9-]+\.){2,}(?:com|net|local|corp|org)\b(?::\d+)?", re.I)
_CRED = re.compile(r"(password|passwd|pwd)\s*[=:]\s*['\"]?[^\s'\";#<>]{3,}|identified\s+by\s+\S+", re.I)


def clean_sql(text):
    text = _COMMENT.sub(" ", text or "")
    text = _STRING.sub("''", text)
    return _NOISE_FROM.sub(" ", text)


def _tab(tok):
    t = norm(re.sub(r"\s+", "", tok))
    if not t or t in SQL_RESERVED or t.startswith("(") or t[0].isdigit():
        return ""
    if short(t) in SQL_RESERVED:
        return ""
    return t


def analyze_sql(text):
    """-> {'reads': set, 'writes': {tab: set(ops)}, 'calls': set(raw idents),
           'vars': set, 'odiref': set, 'dblinks': set, 'jdbc': set, 'hosts': set,
           'creds': int, 'select_star': bool}"""
    res = {"reads": set(), "writes": defaultdict(set), "calls": set(), "vars": set(),
           "odiref": set(), "dblinks": set(), "jdbc": set(), "hosts": set(),
           "creds": 0, "select_star": False}
    if not text or not text.strip():
        return res
    raw = text
    res["vars"] = {norm(v) for v in _ODI_VAR.findall(raw) if not v.isdigit()}
    res["odiref"] = {m[0] for m in _ODIREF.findall(raw)}
    res["jdbc"] = set(_JDBC.findall(raw))
    res["hosts"] = {h.lower() for h in _HOST.findall(raw)}
    res["creds"] = len(_CRED.findall(raw))
    sql = clean_sql(raw)
    res["select_star"] = bool(re.search(r"\bselect\s+(?:/\*.*?\*/\s*)?(?:distinct\s+)?\*\s+from", sql, re.I))
    for op, rx in _WRITES:
        for m in rx.finditer(sql):
            t = _tab(m.group(1))
            if t:
                res["writes"][t].add(op)
    for m in _USING.finditer(sql):
        t = _tab(m.group(1))
        if t:
            res["reads"].add(t)
    # FROM / JOIN with comma lists at paren depth 0
    for m in _FROM_JOIN.finditer(sql):
        i, depth, buf, n = m.end(), 0, [], len(sql)
        while i < n:
            ch = sql[i]
            if ch == "(":
                depth += 1
            elif ch == ")":
                if depth == 0:
                    break
                depth -= 1
            elif ch == ";" and depth == 0:
                break
            elif depth == 0 and ch.isalpha() and (i == 0 or not (sql[i - 1].isalnum() or sql[i - 1] in "_$#.")):
                if _FROM_STOP.match(sql, i):
                    break
            buf.append(ch)
            i += 1
            if len(buf) > 4000:
                break
        seg = "".join(buf)
        if m.group(1).lower() == "join":
            seg = seg.split(",")[0]
        for piece in _split_top(seg):
            piece = piece.strip()
            if not piece or piece.startswith("("):
                continue
            mm = re.match(_IDENT, piece)
            if mm:
                t = _tab(mm.group(0))
                if t:
                    res["reads"].add(t)
    for tabs in (res["reads"], res["writes"].keys()):
        for t in tabs:
            if "@" in t:
                res["dblinks"].add(t.split("@", 1)[1])
    for m in _CALL.finditer(sql):
        res["calls"].add(norm(re.sub(r"\s+", "", m.group(1))))
    for m in _EXEC.finditer(sql):
        res["calls"].add(norm(re.sub(r"\s+", "", m.group(1))))
    return res


def _split_top(seg):
    out, depth, cur = [], 0, []
    for ch in seg:
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth -= 1
        if ch == "," and depth == 0:
            out.append("".join(cur))
            cur = []
        else:
            cur.append(ch)
    out.append("".join(cur))
    return out


def load_pattern(ops_by_table):
    ops = set()
    for s in ops_by_table.values():
        ops |= s
    if not ops:
        return "READ_ONLY / NON-SQL"
    if "MERGE" in ops:
        return "MERGE (upsert)"
    if "TRUNCATE" in ops and "INSERT" in ops:
        return "TRUNCATE + RELOAD"
    if "DELETE" in ops and "INSERT" in ops:
        return "DELETE + INSERT"
    if "UPDATE" in ops and "INSERT" in ops:
        return "INSERT + UPDATE (incremental)"
    if "INSERT" in ops:
        return "APPEND (insert)"
    if "UPDATE" in ops:
        return "UPDATE only"
    if "DELETE" in ops or "TRUNCATE" in ops:
        return "PURGE / DELETE"
    return "DDL"


# ---------------------------------------------------------------- ODI XML
def parse_odi_xml(stream):
    """Flatten a SunopsisExport / VCS xml into [{'_class': 'SnpScen', field: value}]"""
    objs, cur = [], None
    for ev, el in ET.iterparse(stream, events=("start", "end")):
        tag = el.tag.rsplit("}", 1)[-1]
        if ev == "start":
            if tag == "Object":
                cur = {"_class": (el.get("class") or "").rsplit(".", 1)[-1]}
            continue
        if tag == "Field" and cur is not None:
            name = el.get("name")
            if name:
                cur[name] = el.text or ""
        elif tag == "Object" and cur is not None:
            objs.append(cur)
            cur = None
            el.clear()
    return objs


def odi_area(relpath):
    p = relpath.lower()
    for key, area in (("master repository", "MASTER"), ("work repository", "WORK"),
                      ("scenarios and load plans", "RUNTIME"), ("release", "RELEASE")):
        if key in p:
            return area
    return "OTHER"


def name_from_file(fn):
    """SCEN_SCN_X_Version_001.xml -> ('SCN_X','001'); LP_LP_Y.xml -> ('LP_Y','')"""
    base = os.path.splitext(os.path.basename(fn))[0]
    ver = ""
    m = re.search(r"_Version_(\w+)$", base, re.I)
    if m:
        ver, base = m.group(1), base[: m.start()]
    for pre in ("SCEN_", "LP_"):
        if base.upper().startswith(pre):
            base = base[len(pre):]
            break
    return base.upper(), ver


class ODIModel:
    def __init__(self):
        self.scenarios = {}      # (name, ver) -> dict
        self.loadplans = {}      # name -> dict
        self.objects = []        # work/master objects rows
        self.by_class_id = {}    # (class, id) -> obj  (id resolution across files)
        self.raw = []            # (area, relpath, file, objs)
        self.releases = []       # rows
        self.errors = []


ID_FIELDS = {  # class -> internal id field
    "SnpContext": "IContext", "SnpConnect": "IConnect", "SnpPschema": "IPschema",
    "SnpLschema": "ILschema", "SnpTechno": "ITechno", "SnpAgent": "IAgent",
    "SnpLagent": "ILagent", "SnpProject": "IProject", "SnpFolder": "IFolder",
    "SnpModel": "IMod", "SnpSubModel": "ISmod", "SnpTable": "ITable",
    "SnpMapping": "IMapping", "SnpPackage": "IPackage", "SnpTrt": "ITrt",
    "SnpVar": "IVar", "SnpScen": "ScenNo", "SnpLoadPlan": "ILoadPlan",
    "SnpLpStep": "ILpStep",
}


def load_odi(odi_root, git, model):
    files = []
    for path in walk(odi_root):
        ext = os.path.splitext(path)[1].lower()
        if ext == ".xml":
            files.append((path, None))
        elif ext == ".zip":
            try:
                with zipfile.ZipFile(path) as z:
                    for n in z.namelist():
                        if n.lower().endswith(".xml"):
                            files.append((path, n))
            except zipfile.BadZipFile as e:
                model.errors.append({"file": rel(path, odi_root), "error": f"zip: {e}"})
    for path, member in files:
        r = rel(path, odi_root) + (f"!{member}" if member else "")
        try:
            if member:
                with zipfile.ZipFile(path) as z:
                    objs = parse_odi_xml(io.BytesIO(z.read(member)))
            else:
                with open(path, "rb") as f:
                    objs = parse_odi_xml(f)
        except ET.ParseError as e:
            model.errors.append({"file": r, "error": f"xml: {e}"})
            continue
        if not objs:
            continue
        for o in objs:
            idf = ID_FIELDS.get(o["_class"])
            if idf and o.get(idf):
                model.by_class_id[(o["_class"], o[idf])] = o
        model.raw.append((odi_area(r), r, path, member, objs))


def resolve(model, cls, id_):
    return model.by_class_id.get((cls, id_), {}) if id_ else {}


def build_odi(model, git, odi_root):
    master_rows, work_rows, var_rows = [], [], []
    for area, r, path, member, objs in model.raw:
        gi = git.get(path)
        by = defaultdict(list)
        for o in objs:
            by[o["_class"]].append(o)
        release = ""
        if area == "RELEASE":
            parts = r.split("/")
            idx = next((i for i, p in enumerate(parts) if "release" in p.lower()), 0)
            release = parts[idx + 1] if idx + 1 < len(parts) - 1 else parts[idx]

        # ---- scenarios
        for s in by.get("SnpScen", []):
            fname, fver = name_from_file(r.split("!")[-1])
            name = norm(g(s, "ScenName", "Name")) or fname
            ver = g(s, "ScenVersion", "Version") or fver
            scen_no = g(s, "ScenNo", "IScen")
            key = (name, ver)
            rec = {
                "name": name, "version": ver, "scen_no": scen_no,
                "description": g(s, "Description", "ScenDescription"),
                "source_type": g(s, "OrigObjType", "SrcObjType", "OrigObjectType"),
                "source_name": g(s, "OrigObjName", "SrcObjName", "OrigObjectName"),
                "first_date": g(s, "FirstDate"), "first_user": g(s, "FirstUser"),
                "last_date": g(s, "LastDate"), "last_user": g(s, "LastUser"),
                "area": area, "file": r, "abspath": path, "git": gi, "releases": set(),
                "steps": [], "tasks": [], "texts": [], "vars": set(),
            }
            if release:
                rec["releases"].add(release)
            prev = model.scenarios.get(key)
            if prev and prev["area"] == "RUNTIME" and area != "RUNTIME":
                prev["releases"] |= rec["releases"]
                continue  # keep the canonical "Scenarios and Load Plans" copy
            if prev:
                rec["releases"] |= prev["releases"]
            # steps / tasks / texts belonging to this scenario (same file)
            same = lambda o: not scen_no or not o.get("ScenNo") or o.get("ScenNo") == scen_no
            for st in by.get("SnpScenStep", []):
                if same(st):
                    rec["steps"].append({
                        "nno": g(st, "Nno", "StepNo"), "name": g(st, "StepName", "Name"),
                        "type": g(st, "StepType", "Type"),
                        "ok_next": g(st, "NnoOk", "OkNextStep", "NextStepOk"),
                        "ko_next": g(st, "NnoKo", "KoNextStep", "NextStepKo")})
            txt_by_task = defaultdict(list)
            for tx in by.get("SnpScenTxt", []) + by.get("SnpTxt", []) + by.get("SnpTxtHeader", []):
                if same(tx):
                    k = (g(tx, "Nno"), g(tx, "ScenTaskNo", "TaskNo"))
                    order = g(tx, "OrdTxt", "TxtOrd", "Ord", default="0")
                    txt_by_task[k].append((int(order) if str(order).isdigit() else 0,
                                           g(tx, "IndTxtCol", "TxtType", "IndTxt"),
                                           g(tx, "Txt", "FullText", "Text")))
            for t in by.get("SnpScenTask", []):
                if not same(t):
                    continue
                k = (g(t, "Nno"), g(t, "ScenTaskNo", "TaskNo"))
                chunks = sorted(txt_by_task.get(k, []))
                src = "".join(c[2] for c in chunks if str(c[1]).upper() in ("C", "S", "SRC"))
                tgt = "".join(c[2] for c in chunks if str(c[1]).upper() not in ("C", "S", "SRC"))
                src = src or g(t, "ColTxt", "SrcTxt", "SourceCode")
                tgt = tgt or g(t, "DefTxt", "TgtTxt", "TargetCode")
                rec["tasks"].append({
                    "nno": k[0], "task_no": k[1],
                    "name1": g(t, "TaskName1"), "name2": g(t, "TaskName2"),
                    "name3": g(t, "TaskName3"), "type": g(t, "TaskType", "Type"),
                    "src_lschema": g(t, "ColLschemaName", "SrcLschemaName"),
                    "src_techno": g(t, "ColTechIntName", "ColTechnoName"),
                    "tgt_lschema": g(t, "DefLschemaName", "TgtLschemaName"),
                    "tgt_techno": g(t, "DefTechIntName", "DefTechnoName"),
                    "src_code": src, "tgt_code": tgt})
            # every long textual field of every object in the file -> scenario-level analysis
            for o in objs:
                for fk, fv in o.items():
                    if fk != "_class" and isinstance(fv, str) and len(fv) > 15 and \
                            re.search(r"\b(select|insert|update|delete|merge|truncate|begin|call|exec|startscen|odiRef)\b", fv, re.I):
                        rec["texts"].append(fv)
            for v in by.get("SnpVar", []) + by.get("SnpScenVar", []) + by.get("SnpVarScen", []):
                vn = g(v, "VarName", "Name")
                if vn:
                    rec["vars"].add(norm(vn))
            model.scenarios[key] = rec

        # ---- load plans
        for lp in by.get("SnpLoadPlan", []):
            fname, _ = name_from_file(r.split("!")[-1])
            name = norm(g(lp, "LoadPlanName", "Name")) or fname
            if name in model.loadplans and model.loadplans[name]["area"] == "RUNTIME" and area != "RUNTIME":
                if release:
                    model.loadplans[name]["releases"].add(release)
                continue
            lp_id = g(lp, "ILoadPlan")
            steps = {}
            for st in by.get("SnpLpStep", []):
                if lp_id and st.get("ILoadPlan") and st.get("ILoadPlan") != lp_id:
                    continue
                sid = g(st, "ILpStep") or str(len(steps))
                steps[sid] = {
                    "id": sid, "parent": g(st, "ParILpStep", "IParLpStep", "ParentILpStep"),
                    "name": g(st, "LpStepName", "StepName", "Name"),
                    "type": g(st, "LpStepType", "StepType"),
                    "order": g(st, "StepOrder", "OrdStep", "Ord", default="0"),
                    "scen": norm(g(st, "ScenName")), "scen_ver": g(st, "ScenVersion"),
                    "enabled": g(st, "IndEnabled", "Enabled", default="1"),
                    "restart": g(st, "RestartType"), "exception": g(st, "ExceptionName", "ExcepName"),
                    "max_parallel": g(st, "MaxParallelSteps", "MaxParallel"),
                    "scen_tag": g(st, "ScenTag"), "context": g(st, "ContextCode")}
            lpvars = [g(v, "VarName", "Name") for v in by.get("SnpLpVar", []) + by.get("SnpVar", [])]
            model.loadplans[name] = {
                "name": name, "id": lp_id, "description": g(lp, "Description"),
                "last_date": g(lp, "LastDate"), "last_user": g(lp, "LastUser"),
                "area": area, "file": r, "abspath": path, "git": gi,
                "steps": steps, "vars": {norm(v) for v in lpvars if v},
                "releases": {release} if release else set()}

        # ---- release manifest
        if area == "RELEASE":
            for o in objs:
                nm = g(o, "ScenName", "LoadPlanName", "Name", "TrtName", "PackName", "TableName")
                if nm and o["_class"] in ("SnpScen", "SnpLoadPlan", "SnpMapping", "SnpTrt",
                                          "SnpPackage", "SnpTable", "SnpVar"):
                    model.releases.append({"release": release, "object_class": o["_class"],
                                           "object_name": norm(nm), "version": g(o, "ScenVersion"),
                                           "file": r, "commit_date": gi.get("last_date", ""),
                                           "jira": join(gi.get("jira", []))})

        # ---- master & work objects (anything else of interest)
        for o in objs:
            c = o["_class"]
            if c == "SnpContext":
                master_rows.append({"kind": "CONTEXT", "name": g(o, "ContextCode"),
                                    "detail": g(o, "ContextName"), "default": g(o, "IndDefault"), "file": r})
            elif c == "SnpTechno":
                master_rows.append({"kind": "TECHNOLOGY", "name": g(o, "TechIntName", "TechnoName"),
                                    "detail": g(o, "TechName", "Name"), "file": r})
            elif c == "SnpConnect":
                tech = resolve(model, "SnpTechno", o.get("ITechno"))
                master_rows.append({"kind": "DATA_SERVER", "name": g(o, "ConName", "Name"),
                                    "technology": g(tech, "TechIntName") or g(o, "TechIntName"),
                                    "url": g(o, "JavaUrl", "Url", "JdbcUrl"),
                                    "driver": g(o, "JavaDriver"), "user": g(o, "UserName", "User"),
                                    "host": g(o, "HostName", "ServerName", "DServer"), "file": r})
            elif c == "SnpPschema":
                con = resolve(model, "SnpConnect", o.get("IConnect"))
                master_rows.append({"kind": "PHYSICAL_SCHEMA",
                                    "name": g(o, "ExtName", "PschemaName") or
                                    f"{g(con, 'ConName')}.{g(o, 'SchemaName')}",
                                    "schema": g(o, "SchemaName"), "work_schema": g(o, "WschemaName"),
                                    "data_server": g(con, "ConName"), "file": r})
            elif c == "SnpLschema":
                tech = resolve(model, "SnpTechno", o.get("ITechno"))
                master_rows.append({"kind": "LOGICAL_SCHEMA", "name": g(o, "LschemaName", "Name"),
                                    "technology": g(tech, "TechIntName"), "file": r})
            elif c == "SnpPschemaCont":
                ctx = resolve(model, "SnpContext", o.get("IContext"))
                ls = resolve(model, "SnpLschema", o.get("ILschema"))
                ps = resolve(model, "SnpPschema", o.get("IPschema"))
                con = resolve(model, "SnpConnect", ps.get("IConnect"))
                master_rows.append({"kind": "CONTEXT_MAPPING",
                                    "name": g(ls, "LschemaName") or o.get("ILschema", ""),
                                    "context": g(ctx, "ContextCode") or g(o, "ContextCode"),
                                    "schema": g(ps, "SchemaName"), "work_schema": g(ps, "WschemaName"),
                                    "data_server": g(con, "ConName"), "url": g(con, "JavaUrl"), "file": r})
            elif c in ("SnpAgent", "SnpLagent"):
                master_rows.append({"kind": "AGENT" if c == "SnpAgent" else "LOGICAL_AGENT",
                                    "name": g(o, "AgentName", "LagentName", "Name"),
                                    "host": g(o, "HostName", "ServerName", "AgentHost"),
                                    "port": g(o, "AgentPort", "Port"), "file": r})
            elif c in ("SnpProject", "SnpFolder", "SnpMapping", "SnpPackage", "SnpTrt",
                       "SnpModel", "SnpSubModel", "SnpTable", "SnpSequence", "SnpUfunc"):
                kind = {"SnpProject": "PROJECT", "SnpFolder": "FOLDER", "SnpMapping": "MAPPING",
                        "SnpPackage": "PACKAGE", "SnpModel": "MODEL", "SnpSubModel": "SUB_MODEL",
                        "SnpTable": "DATASTORE", "SnpSequence": "ODI_SEQUENCE",
                        "SnpUfunc": "USER_FUNCTION"}.get(c)
                if c == "SnpTrt":
                    tt = g(o, "TrtType")
                    kind = "PROCEDURE" if tt in ("U", "") else f"KM({tt})"
                folder = resolve(model, "SnpFolder", o.get("IFolder"))
                proj = resolve(model, "SnpProject", o.get("IProject") or folder.get("IProject"))
                mod = resolve(model, "SnpModel", o.get("IMod"))
                work_rows.append({
                    "kind": kind, "name": g(o, "Name", "TrtName", "PackName", "ModName", "TableName",
                                             "FolderName", "ProjectName", "SmodName", "SeqName", "UfuncName"),
                    "code": g(o, "ProjectCode", "CodName", "ResName"),
                    "folder": g(folder, "FolderName"), "project": g(proj, "ProjectName", "ProjectCode"),
                    "model": g(mod, "ModName"), "lschema": g(o, "LschemaName") or g(mod, "LschemaName"),
                    "technology": g(o, "TechIntName") or g(mod, "TechIntName"),
                    "description": g(o, "Description")[:500], "area": area, "file": r,
                    "last_date": g(o, "LastDate"), "last_user": g(o, "LastUser")})
            elif c == "SnpVar" and area in ("WORK", "MASTER", "OTHER"):
                proj = resolve(model, "SnpProject", o.get("IProject"))
                var_rows.append({"name": g(o, "VarName", "Name"),
                                 "project": g(proj, "ProjectCode", "ProjectName") or "GLOBAL",
                                 "datatype": g(o, "VarDatatype", "DataType"),
                                 "default": g(o, "DefV", "DefaultValue")[:200],
                                 "refresh_lschema": g(o, "LschemaName"),
                                 "refresh_sql": g(o, "TxtRefresh", "RefreshQuery")[:1000], "file": r})
    return master_rows, work_rows, var_rows


# ---------------------------------------------------------------- DB scripts
_CREATE = re.compile(
    r"\bcreate\s+(?:or\s+replace\s+)?(?:(?:non)?editionable\s+)?(?:force\s+)?"
    r"(?:global\s+temporary\s+)?(unique\s+|bitmap\s+)?"
    r"(table|materialized\s+view|view|procedure|function|package\s+body|package|sequence|"
    r"index|trigger|public\s+synonym|synonym|type\s+body|type)\s+(" + _IDENT + ")", re.I)
_FK = re.compile(r"alter\s+table\s+(" + _IDENT + r")\s+add\s+(?:constraint\s+\w+\s+)?foreign\s+key"
                 r"\s*\([^)]*\)\s*references\s+(" + _IDENT + ")", re.I)
_INLINE_FK = re.compile(r"\breferences\s+(" + _IDENT + ")", re.I)
FOLDER_TYPE = {"functions": "FUNCTION", "function": "FUNCTION", "index": "INDEX", "indexes": "INDEX",
               "procedure": "PROCEDURE", "procedures": "PROCEDURE", "sequence": "SEQUENCE",
               "sequences": "SEQUENCE", "tables": "TABLE", "table": "TABLE", "views": "VIEW",
               "view": "VIEW", "packages": "PACKAGE", "package": "PACKAGE", "triggers": "TRIGGER",
               "synonyms": "SYNONYM", "types": "TYPE", "mviews": "MATERIALIZED VIEW"}


def schema_from_path(relpath):
    parts = relpath.split("/")[:-1]
    cands = [p for p in parts if p.lower() not in FOLDER_TYPE and p.lower() != "database scripts"]
    for p in reversed(cands):
        if re.fullmatch(r"[A-Z][A-Z0-9_$#]{1,29}", p):
            return p
    return cands[-1].upper() if cands else ""


def load_db_scripts(roots, git):
    objs, fks, errors = [], [], []
    for root in roots:
        for path in walk(root, {".sql", ".pks", ".pkb", ".pck", ".prc", ".fnc", ".vw", ".trg",
                                ".tab", ".seq", ".idx", ".ddl", ".plsql", ".txt", ""}):
            r = rel(path, root)
            text = read_text(path)
            if not text:
                continue
            folder_type = next((FOLDER_TYPE[p.lower()] for p in reversed(r.split("/")[:-1])
                                if p.lower() in FOLDER_TYPE), "")
            path_schema = schema_from_path(r)
            gi = git.get(path)
            ms = list(_CREATE.finditer(text))
            if not ms:
                if folder_type:
                    objs.append(_db_row(folder_type, path_schema,
                                        norm(os.path.splitext(os.path.basename(path))[0]),
                                        "", root, r, gi, {}, note="no CREATE statement found"))
                continue
            for i, m in enumerate(ms):
                body = text[m.start(): ms[i + 1].start() if i + 1 < len(ms) else len(text)]
                otype = re.sub(r"\s+", " ", m.group(2).upper())
                if m.group(1) and otype == "INDEX":
                    otype = f"{m.group(1).strip().upper()} INDEX"
                full = norm(re.sub(r"\s+", "", m.group(3)))
                schema, name = (full.split(".", 1) if "." in full else (path_schema, full))
                extra = {}
                if "INDEX" in otype:
                    mo = re.search(r"\bon\s+(" + _IDENT + r")\s*\(([^)]*)\)", body, re.I)
                    if mo:
                        extra["on_table"] = norm(re.sub(r"\s+", "", mo.group(1)))
                        extra["columns"] = norm(mo.group(2)).replace(" ", "")
                elif otype == "TRIGGER":
                    mo = re.search(r"\bon\s+(" + _IDENT + ")", body, re.I)
                    if mo:
                        extra["on_table"] = norm(re.sub(r"\s+", "", mo.group(1)))
                elif "SYNONYM" in otype:
                    mo = re.search(r"\bfor\s+(" + _IDENT + ")", body, re.I)
                    if mo:
                        extra["on_table"] = norm(re.sub(r"\s+", "", mo.group(1)))
                elif otype == "TABLE":
                    extra["columns"] = _count_columns(body)
                    extra["partitioned"] = "Y" if re.search(r"\bpartition\s+by\b", body, re.I) else "N"
                    mt = re.search(r"\btablespace\s+(\w+)", body, re.I)
                    extra["tablespace"] = mt.group(1).upper() if mt else ""
                    for mf in _INLINE_FK.finditer(body):
                        fks.append({"table": f"{schema}.{name}".strip("."),
                                    "references": norm(re.sub(r"\s+", "", mf.group(1))), "file": r})
                if otype not in ("TABLE", "SEQUENCE") and "INDEX" not in otype and "SYNONYM" not in otype:
                    extra["sql"] = analyze_sql(body)
                objs.append(_db_row(otype, schema, name, folder_type, root, r, gi, extra))
            for mf in _FK.finditer(text):
                fks.append({"table": norm(re.sub(r"\s+", "", mf.group(1))),
                            "references": norm(re.sub(r"\s+", "", mf.group(2))), "file": r})
    return objs, fks, errors


def _count_columns(body):
    m = re.search(r"\(", body)
    if not m:
        return ""
    depth, i, start = 0, m.start(), m.start() + 1
    for i in range(m.start(), len(body)):
        if body[i] == "(":
            depth += 1
        elif body[i] == ")":
            depth -= 1
            if depth == 0:
                break
    cols = [c.strip() for c in _split_top(clean_sql(body[start:i]))]
    cols = [c for c in cols if c and not re.match(r"(constraint|primary|foreign|unique|check|supplemental)\b", c, re.I)]
    return str(len(cols))


def _db_row(otype, schema, name, folder_type, root, r, gi, extra, note=""):
    return {"object_type": otype, "schema": schema, "name": name, "folder_type": folder_type,
            "repo": os.path.basename(os.path.abspath(root)), "file": r,
            "on_table": extra.get("on_table", ""), "columns": extra.get("columns", ""),
            "partitioned": extra.get("partitioned", ""), "tablespace": extra.get("tablespace", ""),
            "sql": extra.get("sql"), "note": note,
            "last_commit_date": gi.get("last_date", ""), "last_author": gi.get("last_author", ""),
            "jira": join(gi.get("jira", []))}


# ---------------------------------------------------------------- Autosys JIL
_JIL_ATTR = re.compile(r"(?:^|\s)([a-z_]+):\s*")


def load_autosys(roots, git):
    jobs = []
    for root in roots:
        for path in walk(root):
            text = read_text(path)
            if "insert_job" not in text and "update_job" not in text:
                continue
            r = rel(path, root)
            cur = None
            for line in text.splitlines():
                s = line.strip()
                if not s or s.startswith("/*") or s.startswith("#"):
                    continue
                if s.lower().startswith("command:"):
                    if cur is not None:
                        cur["command"] = s.split(":", 1)[1].strip()
                    continue
                keys = list(_JIL_ATTR.finditer(s))
                if not keys or keys[0].start() > 0:
                    continue
                for i, k in enumerate(keys):
                    key = k.group(1).lower()
                    val = s[k.end(): keys[i + 1].start() if i + 1 < len(keys) else len(s)].strip().strip('"')
                    if key in ("insert_job", "update_job"):
                        cur = {"job": val, "file": r, "abspath": path, "repo": os.path.basename(root)}
                        jobs.append(cur)
                    elif cur is not None:
                        cur[key] = val
    for j in jobs:
        gi = git.get(j.pop("abspath"))
        j["jira"] = join(gi.get("jira", []))
        j["last_commit_date"] = gi.get("last_date", "")
    return jobs


def jil_deps(condition):
    return re.findall(r"\b(?:s|success|d|done|f|failure|n|notrunning|t|terminated|e|exitcode)\(\s*([\w.#-]+)",
                      condition or "", re.I)


# ---------------------------------------------------------------- generic scripts
_STARTSCEN = re.compile(r"startscen(?:\.sh|\.bat|\.cmd)?\s+[\"']?([\w-]+)[\"']?\s+[\"']?(\w+|-1)?[\"']?\s*[\"']?(\w+)?", re.I)
_STARTLP = re.compile(r"startloadplan(?:\.sh|\.bat|\.cmd)?\s+[\"']?([\w-]+)[\"']?\s*[\"']?(\w+)?", re.I)
_OTHER_EXT = TEXT_EXT - {".xml", ".json"}


def scan_scripts(roots, git, scen_names, lp_names, db_callables):
    rows = []
    token_re = re.compile(r"[A-Za-z_][\w$#]{3,}")
    for root in roots:
        for path in walk(root, _OTHER_EXT):
            text = read_text(path, 5_000_000)
            if not text:
                continue
            r = rel(path, root)
            gi = git.get(path)
            hits = set()
            for m in _STARTSCEN.finditer(text):
                hits.add(("RUNS_SCENARIO", norm(m.group(1)), m.group(2) or ""))
            for m in _STARTLP.finditer(text):
                hits.add(("RUNS_LOAD_PLAN", norm(m.group(1)), ""))
            toks = {t.upper() for t in token_re.findall(text)}
            for t in toks & scen_names:
                hits.add(("REFERENCES_SCENARIO", t, ""))
            for t in toks & lp_names:
                hits.add(("REFERENCES_LOAD_PLAN", t, ""))
            for t in toks & db_callables:
                hits.add(("REFERENCES_DB_OBJECT", t, ""))
            for kind, target, ver in sorted(hits):
                rows.append({"repo": os.path.basename(os.path.abspath(root)), "file": r,
                             "relation": kind, "target": target, "version": ver,
                             "last_commit_date": gi.get("last_date", ""),
                             "jira": join(gi.get("jira", []))})
    return rows


# ---------------------------------------------------------------- 360 assembly
def build_360(args, model, master_rows, work_rows, var_rows, db_objs, fks, jobs, script_rows):
    out = defaultdict(list)
    catalog = defaultdict(list)  # SHORT NAME -> db rows
    for d in db_objs:
        catalog[d["name"]].append(d)
    callable_types = ("PROCEDURE", "FUNCTION", "PACKAGE", "PACKAGE BODY")
    callables = {n for n, rows in catalog.items() if any(x["object_type"] in callable_types for x in rows)}
    table_like = {n for n, rows in catalog.items()
                  if any(x["object_type"] in ("TABLE", "VIEW", "MATERIALIZED VIEW", "SYNONYM") for x in rows)}

    def classify_call(ident):
        parts = ident.split(".")
        for p in (parts[-1], parts[0]):
            if p in callables:
                types = {x["object_type"] for x in catalog[p]}
                kind = "FUNCTION" if types == {"FUNCTION"} else ("PACKAGE" if "PACKAGE" in " ".join(types) else "PROCEDURE")
                return kind, ident
        return None, None

    edges = []

    def edge(st, s, rel_, dt, d, detail=""):
        edges.append({"src_type": st, "src": s, "relation": rel_, "dst_type": dt, "dst": d, "detail": detail})

    table_use = defaultdict(lambda: {"readers": set(), "writers": set(), "ops": set()})

    # ---- Autosys lookup
    job_by_scen, job_by_lp = defaultdict(list), defaultdict(list)
    scen_names = {k[0] for k in model.scenarios}
    lp_names = set(model.loadplans)
    for j in jobs:
        cmd = j.get("command", "")
        targets = set()
        for m in _STARTSCEN.finditer(cmd):
            targets.add(("S", norm(m.group(1))))
            j["odi_context"] = m.group(3) or j.get("odi_context", "")
        for m in _STARTLP.finditer(cmd):
            targets.add(("L", norm(m.group(1))))
            j["odi_context"] = m.group(2) or j.get("odi_context", "")
        toks = {t.upper() for t in re.findall(r"[A-Za-z_][\w$#-]{3,}", cmd)}
        targets |= {("S", t) for t in toks & scen_names} | {("L", t) for t in toks & lp_names}
        j["odi_targets"] = join(f"{'SCENARIO' if k == 'S' else 'LOAD_PLAN'}:{n}" for k, n in targets)
        for k, n in targets:
            (job_by_scen if k == "S" else job_by_lp)[n].append(j)
            edge("AUTOSYS_JOB", j["job"], "RUNS", "SCENARIO" if k == "S" else "LOAD_PLAN", n)
        if j.get("box_name"):
            edge("AUTOSYS_BOX", j["box_name"], "CONTAINS", "AUTOSYS_JOB", j["job"])
        for dep in jil_deps(j.get("condition", "")):
            edge("AUTOSYS_JOB", j["job"], "DEPENDS_ON", "AUTOSYS_JOB", dep, j.get("condition", ""))
    box_of = {j["job"]: j.get("box_name", "") for j in jobs}
    job_rec = {j["job"]: j for j in jobs}

    def schedule(j):
        b = job_rec.get(j.get("box_name", ""), {})
        parts = []
        for src in (j, b):
            for k in ("date_conditions", "days_of_week", "run_calendar", "start_times",
                      "start_mins", "run_window", "timezone"):
                if src.get(k) and src.get(k) != "0" and k != "date_conditions":
                    parts.append(f"{k}={src[k]}")
            if parts:
                break
        return " ".join(parts)

    # ---- load plan steps
    lp_step_rows, scen_in_lp = [], defaultdict(list)
    for lp in model.loadplans.values():
        steps = lp["steps"]

        def path_of(s, seen=None):
            seen = seen or set()
            if s["id"] in seen:
                return s["name"]
            seen.add(s["id"])
            p = steps.get(s["parent"])
            return (path_of(p, seen) + " > " if p else "") + (s["name"] or s["type"])

        def sort_key(s, depth=0):
            p = steps.get(s["parent"])
            own = (int(s["order"]) if str(s["order"]).lstrip("-").isdigit() else 0, s["id"])
            return (sort_key(p, depth + 1) if p and depth < 50 else ()) + (own,)

        for s in sorted(steps.values(), key=sort_key):  # depth-first, in step order
            pth = path_of(s)
            enabled = "N" if str(s["enabled"]).strip() in ("0", "false", "N") else "Y"
            lp_step_rows.append({"load_plan": lp["name"], "step_id": s["id"], "parent_id": s["parent"],
                                 "order": s["order"], "step_name": s["name"], "step_type": s["type"],
                                 "path": pth, "scenario": s["scen"], "scenario_version": s["scen_ver"],
                                 "enabled": enabled, "restart_type": s["restart"],
                                 "exception_step": s["exception"], "max_parallel": s["max_parallel"],
                                 "context": s["context"]})
            if s["scen"]:
                scen_in_lp[s["scen"]].append((lp["name"], pth, enabled))
                edge("LOAD_PLAN", lp["name"], "RUNS", "SCENARIO", s["scen"],
                     f"step={pth}; version={s['scen_ver'] or 'latest'}; enabled={enabled}")
    out["loadplan_steps"] = lp_step_rows

    # ---- work objects -> scenario generated-from
    work_by_name = defaultdict(list)
    for w in work_rows:
        if w["kind"] in ("MAPPING", "PACKAGE", "PROCEDURE"):
            work_by_name[norm(w["name"])].append(w)

    scripts_by_target = defaultdict(list)
    for x in script_rows:
        scripts_by_target[x["target"]].append(x)

    # ---- scenarios
    scen_rows, task_rows, usage_rows, findings = [], [], [], []
    scen_facts = {}
    versions = Counter(k[0] for k in model.scenarios)
    for (name, ver), s in sorted(model.scenarios.items()):
        agg = analyze_sql("\n;\n".join(s["texts"] + [t["src_code"] + "\n;\n" + t["tgt_code"] for t in s["tasks"]]))
        lschemas, technos, kms = set(), set(), set()
        for t in s["tasks"]:
            lschemas |= {x for x in (t["src_lschema"], t["tgt_lschema"]) if x}
            technos |= {x for x in (t["src_techno"], t["tgt_techno"]) if x}
            for nm in (t["name1"], t["name2"], t["name3"]):
                if re.search(r"\b[LICJRS]KM\b", nm or ""):
                    kms.add(nm)
            ta = analyze_sql(t["src_code"] + "\n;\n" + t["tgt_code"])
            task_rows.append({"scenario": name, "version": ver, "step_no": t["nno"], "task_no": t["task_no"],
                              "task_name1": t["name1"], "task_name2": t["name2"], "task_name3": t["name3"],
                              "task_type": t["type"], "src_lschema": t["src_lschema"],
                              "src_technology": t["src_techno"], "tgt_lschema": t["tgt_lschema"],
                              "tgt_technology": t["tgt_techno"],
                              "reads": join(ta["reads"]), "writes": join(f"{k}({'/'.join(sorted(v))})" for k, v in ta["writes"].items()),
                              "code_chars": len(t["src_code"]) + len(t["tgt_code"]),
                              "code_preview": (t["tgt_code"] or t["src_code"])[:1500]})
        reads = {t for t in agg["reads"] if not ODI_WORK_TABLE.match(short(t))}
        writes = {t: ops for t, ops in agg["writes"].items() if not ODI_WORK_TABLE.match(short(t))}
        reads -= set(writes)  # DELETE FROM / UPDATE targets are not reads
        work_tables = {short(t) for t in list(agg["reads"]) + list(agg["writes"]) if ODI_WORK_TABLE.match(short(t))}
        procs, funcs = set(), set()
        for c in agg["calls"]:
            kind, ident = classify_call(c)
            if kind == "FUNCTION":
                funcs.add(ident)
            elif kind:
                procs.add(ident)
        for t in reads:
            table_use[short(t)]["readers"].add(name)
            usage_rows.append({"scenario": name, "version": ver, "table": t, "operation": "READ",
                               "in_db_scripts": "Y" if short(t) in table_like else "N"})
            edge("SCENARIO", name, "READS", "TABLE", t)
        for t, ops in writes.items():
            table_use[short(t)]["writers"].add(name)
            table_use[short(t)]["ops"] |= ops
            usage_rows.append({"scenario": name, "version": ver, "table": t, "operation": "/".join(sorted(ops)),
                               "in_db_scripts": "Y" if short(t) in table_like else "N"})
            edge("SCENARIO", name, "WRITES", "TABLE", t, "/".join(sorted(ops)))
        for p in procs:
            edge("SCENARIO", name, "CALLS", "DB_PROCEDURE", p)
        for f in funcs:
            edge("SCENARIO", name, "CALLS", "DB_FUNCTION", f)
        for ls in lschemas:
            edge("SCENARIO", name, "USES_LOGICAL_SCHEMA", "LOGICAL_SCHEMA", ls)
        src_objs = work_by_name.get(norm(s["source_name"])) or work_by_name.get(name) or \
            work_by_name.get(re.sub(r"^SCN_", "", name)) or []
        src_obj = src_objs[0] if src_objs else {}
        if src_obj:
            edge("SCENARIO", name, "GENERATED_FROM", src_obj["kind"], src_obj["name"], src_obj.get("folder", ""))
        gi = s["git"]
        lps = scen_in_lp.get(name, [])
        sj = job_by_scen.get(name, [])
        lp_jobs = [j for lp_ in {x[0] for x in lps} for j in job_by_lp.get(lp_, [])]
        sr = scripts_by_target.get(name, [])
        facts = {
            "name": name, "version": ver, "reads": reads, "writes": writes, "procs": procs, "funcs": funcs,
            "lschemas": lschemas, "technos": technos, "kms": kms, "vars": s["vars"] | agg["vars"],
            "dblinks": agg["dblinks"], "work_tables": work_tables, "lps": lps,
            "jobs": sj + lp_jobs, "scripts": sr, "src_obj": src_obj, "s": s, "agg": agg}
        scen_facts[(name, ver)] = facts
        scen_rows.append({
            "scenario": name, "version": ver, "description": s["description"][:1000],
            "source_object_type": s["source_type"] or src_obj.get("kind", ""),
            "source_object_name": s["source_name"] or src_obj.get("name", ""),
            "source_folder": src_obj.get("folder", ""), "project": src_obj.get("project", ""),
            "repo_area": s["area"], "file": s["file"], "step_count": len(s["steps"]),
            "task_count": len(s["tasks"]),
            "steps": join((f"{x['nno']}:{x['name']}[{x['type']}]" for x in s["steps"]), " | "),
            "load_plans": join(x[0] for x in lps),
            "autosys_jobs": join(j["job"] for j in sj + lp_jobs),
            "scripts": join(f"{x['repo']}/{x['file']}" for x in sr),
            "reads": join(reads), "writes": join(f"{k}({'/'.join(sorted(v))})" for k, v in writes.items()),
            "db_procedures": join(procs), "db_functions": join(funcs), "db_links": join(agg["dblinks"]),
            "logical_schemas": join(lschemas), "technologies": join(technos), "kms": join(kms),
            "variables": join(facts["vars"]), "odi_work_tables": join(work_tables),
            "load_pattern": load_pattern(writes), "odi_last_date": s["last_date"],
            "odi_last_user": s["last_user"], "releases": join(s["releases"]),
            "git_last_date": gi.get("last_date", ""), "git_last_author": gi.get("last_author", ""),
            "git_last_subject": gi.get("last_subject", ""), "git_commits": gi.get("n_commits", ""),
            "jira": join(gi.get("jira", []))})

        # findings
        def f(sev, code, msg):
            findings.append({"severity": sev, "code": code, "object_type": "SCENARIO",
                             "object": f"{name} v{ver}", "detail": msg, "file": s["file"]})
        if not lps and not sj and not sr:
            f("MEDIUM", "ORPHAN_SCENARIO", "Not referenced by any load plan, Autosys job or script")
        if versions[name] > 1:
            f("LOW", "MULTIPLE_VERSIONS", f"{versions[name]} versions exported for {name}")
        if agg["creds"]:
            f("HIGH", "HARDCODED_CREDENTIAL", f"{agg['creds']} password-like literal(s) in scenario code")
        if agg["jdbc"] or agg["hosts"]:
            f("MEDIUM", "HARDCODED_ENDPOINT", join(agg["jdbc"] | agg["hosts"]))
        if agg["select_star"]:
            f("LOW", "SELECT_STAR", "SELECT * used - schema drift risk")
        missing = [t for t in list(reads) + list(writes) if short(t) not in table_like and "@" not in t]
        if missing and db_objs:
            f("INFO", "TABLE_NOT_IN_DB_SCRIPTS", join(missing))
        if not gi.get("jira"):
            f("LOW", "NO_JIRA_TRACE", f"Last commit: {gi.get('last_subject', 'n/a')}")
        if not reads and not writes and s["tasks"]:
            f("INFO", "NO_SQL_LINEAGE", "No table lineage resolved (OS/Jython/tool steps or unresolved odiRef)")
    out["scenarios"], out["scenario_tasks"], out["table_usage"] = scen_rows, task_rows, usage_rows

    # ---- scenario data dependencies (A writes T, B reads T -> B after A)
    deps = []
    for t, u in table_use.items():
        for w in u["writers"]:
            for r_ in u["readers"]:
                if w != r_:
                    deps.append({"upstream_scenario": w, "downstream_scenario": r_, "via_table": t})
                    edge("SCENARIO", r_, "DEPENDS_ON_DATA_FROM", "SCENARIO", w, t)
    out["scenario_dependencies"] = deps
    up, down = defaultdict(set), defaultdict(set)
    for d in deps:
        up[d["downstream_scenario"]].add(d["upstream_scenario"])
        down[d["upstream_scenario"]].add(d["downstream_scenario"])

    # ---- load plans
    lp_rows = []
    for lp in model.loadplans.values():
        gi = lp["git"]
        scs = [s["scen"] for s in lp["steps"].values() if s["scen"]]
        js = job_by_lp.get(lp["name"], [])
        lp_rows.append({"load_plan": lp["name"], "description": lp["description"][:1000],
                        "step_count": len(lp["steps"]), "scenario_count": len(set(scs)),
                        "scenarios": join(scs),
                        "disabled_steps": sum(1 for s in lp["steps"].values()
                                              if str(s["enabled"]).strip() in ("0", "false", "N")),
                        "variables": join(lp["vars"]), "autosys_jobs": join(j["job"] for j in js),
                        "schedule": join(schedule(j) for j in js), "odi_last_date": lp["last_date"],
                        "odi_last_user": lp["last_user"], "file": lp["file"], "releases": join(lp["releases"]),
                        "git_last_date": gi.get("last_date", ""), "jira": join(gi.get("jira", []))})
        for s in lp["steps"].values():
            if s["scen"] and not any(k[0] == s["scen"] for k in model.scenarios):
                findings.append({"severity": "HIGH", "code": "LP_MISSING_SCENARIO", "object_type": "LOAD_PLAN",
                                 "object": lp["name"], "detail": f"Step '{s['name']}' runs {s['scen']} which is not in the repo",
                                 "file": lp["file"]})
            if str(s["enabled"]).strip() in ("0", "false", "N"):
                findings.append({"severity": "LOW", "code": "LP_STEP_DISABLED", "object_type": "LOAD_PLAN",
                                 "object": lp["name"], "detail": s["name"], "file": lp["file"]})
        if not js and not scripts_by_target.get(lp["name"]):
            findings.append({"severity": "MEDIUM", "code": "UNSCHEDULED_LOAD_PLAN", "object_type": "LOAD_PLAN",
                             "object": lp["name"], "detail": "No Autosys job / script starts this load plan",
                             "file": lp["file"]})
    out["loadplans"] = lp_rows

    # ---- DB objects + their own lineage
    db_rows = []
    used_by = defaultdict(set)
    for f_ in scen_facts.values():
        for t in list(f_["reads"]) + list(f_["writes"]):
            used_by[short(t)].add(f_["name"])
        for c in f_["procs"] | f_["funcs"]:
            for p in c.split("."):
                used_by[p].add(f_["name"])
    for x in script_rows:
        if x["relation"] == "REFERENCES_DB_OBJECT":
            used_by[x["target"]].add(f"script:{x['file']}")
    for d in db_objs:
        sq = d.pop("sql", None) or {}
        fq = f"{d['schema']}.{d['name']}".strip(".")
        rd = {t for t in sq.get("reads", set()) if short(t) != d["name"]}
        wr = {t for t in sq.get("writes", {}) if short(t) != d["name"]}
        calls = set()
        for c in sq.get("calls", set()):
            k, ident = classify_call(c)
            if k and short(ident) != d["name"]:
                calls.add(ident)
        for t in rd:
            edge(d["object_type"], fq, "READS", "TABLE", t)
            used_by[short(t)].add(fq)
        for t in wr:
            edge(d["object_type"], fq, "WRITES", "TABLE", t)
            used_by[short(t)].add(fq)
        for c in calls:
            edge(d["object_type"], fq, "CALLS", "DB_CALLABLE", c)
            for p in c.split("."):
                used_by[p].add(fq)
        if d["on_table"]:
            edge(d["object_type"], fq, "ON", "TABLE", d["on_table"])
        d.update({"reads": join(rd), "writes": join(wr), "calls": join(calls)})
        db_rows.append(d)
    for d in db_rows:
        u = used_by.get(d["name"], set()) - {f"{d['schema']}.{d['name']}".strip(".")}
        d["used_by"] = join(u)
        d["used_by_count"] = len(u)
        if not u and d["object_type"] in ("TABLE", "VIEW", "PROCEDURE", "FUNCTION", "PACKAGE", "MATERIALIZED VIEW"):
            findings.append({"severity": "INFO", "code": "DB_OBJECT_UNREFERENCED", "object_type": d["object_type"],
                             "object": f"{d['schema']}.{d['name']}", "file": d["file"],
                             "detail": "Not referenced by any scenario, DB object or script in scope"})
    for fk in fks:
        edge("TABLE", fk["table"], "FK_REFERENCES", "TABLE", fk["references"])
    out["db_objects"], out["db_foreign_keys"] = db_rows, fks

    # ---- table 360
    t360 = []
    for t, u in sorted(table_use.items()):
        cat = catalog.get(t, [])
        t360.append({"table": t, "db_object_type": join(x["object_type"] for x in cat),
                     "schema": join(x["schema"] for x in cat), "in_db_scripts": "Y" if cat else "N",
                     "written_by_scenarios": join(u["writers"]), "read_by_scenarios": join(u["readers"]),
                     "write_operations": join(u["ops"]),
                     "db_objects_using": join(x for x in used_by.get(t, set()) if "." in x and not x.startswith("script:")),
                     "indexes": join(f"{x['name']}({x['columns']})" for x in db_rows
                                     if "INDEX" in x["object_type"] and short(x["on_table"]) == t),
                     "layer": first_match(LAYER_RULES, t)})
    out["table_360"] = t360

    # ---- autosys, scripts, master, work, vars, releases
    for j in jobs:
        j["box_name"] = j.get("box_name", "")
        j["schedule"] = schedule(j)
        j["depends_on"] = join(jil_deps(j.get("condition", "")))
    out["autosys_jobs"] = jobs
    out["script_references"] = script_rows
    out["odi_topology"] = master_rows
    out["odi_work_objects"] = work_rows
    out["odi_variables"] = var_rows
    out["releases"] = model.releases
    for lm in master_rows:
        if lm["kind"] == "CONTEXT_MAPPING":
            edge("LOGICAL_SCHEMA", lm["name"], "MAPS_TO", "PHYSICAL_SCHEMA",
                 f"{lm.get('data_server', '')}.{lm.get('schema', '')}", f"context={lm.get('context', '')}")
    for rr in model.releases:
        edge("RELEASE", rr["release"], "CONTAINS", rr["object_class"], rr["object_name"])
    for x in script_rows:
        edge("SCRIPT", f"{x['repo']}/{x['file']}", x["relation"], "", x["target"])
    out["lineage_edges"] = edges
    out["findings"] = sorted(findings, key=lambda x: ["HIGH", "MEDIUM", "LOW", "INFO"].index(x["severity"]))
    out["parse_errors"] = model.errors

    # ---- CP360 ODI_360 sheet
    flags_by_obj = defaultdict(set)
    for x in findings:
        flags_by_obj[x["object"]].add(x["code"])
    cp = []
    order = 0
    lp_list = sorted(model.loadplans.values(), key=lambda x: x["name"])
    for lp in lp_list:
        order += 1
        js = job_by_lp.get(lp["name"], [])
        scs = {s["scen"] for s in lp["steps"].values() if s["scen"]}
        tg = set()
        for (n, v), fx in scen_facts.items():
            if n in scs:
                tg |= set(fx["writes"])
        lp_find = flags_by_obj[lp["name"]]
        cp.append(_cp_row(args, order, "LOAD_PLAN", lp["name"], "", lp["description"],
                          domain=first_match(DOMAIN_RULES, lp["name"]),
                          layer="Orchestration", load_plans=lp["name"], lp_path="(root)",
                          jobs=js, box=join(j.get("box_name", "") for j in js),
                          sched=join(schedule(j) for j in js), tgt=tg, scen_in=scs,
                          flags=lp_find, git=lp["git"], odi_date=lp["last_date"],
                          odi_user=lp["last_user"], path=lp["file"], releases=lp["releases"],
                          criticality="HIGH" if js else "MEDIUM", component="ODI_LOAD_PLAN"))
    for (name, ver), fx in sorted(scen_facts.items()):
        order += 1
        s = fx["s"]
        flags = flags_by_obj[f"{name} v{ver}"]
        crit = "HIGH" if (fx["jobs"] and len(down[name]) >= 1) or any("INTRADAY" in x[0] for x in fx["lps"]) \
            else ("MEDIUM" if fx["jobs"] or fx["lps"] else "LOW")
        tgt_names = [short(t) for t in fx["writes"]]
        cp.append(_cp_row(
            args, order, "SCENARIO", name, ver, s["description"],
            domain=first_match(DOMAIN_RULES, name, s["source_name"], *tgt_names),
            layer=first_match(LAYER_RULES, *tgt_names, name) or "Unclassified",
            src_type=s["source_type"] or fx["src_obj"].get("kind", ""),
            src_name=s["source_name"] or fx["src_obj"].get("name", ""),
            load_plans=join(x[0] for x in fx["lps"]), lp_path=join(x[1] for x in fx["lps"]),
            jobs=fx["jobs"], box=join(box_of.get(j["job"], "") for j in fx["jobs"]),
            sched=join(schedule(j) for j in fx["jobs"]), up=up[name], down=down[name],
            src=fx["reads"], tgt=set(fx["writes"]), dblinks=fx["dblinks"], procs=fx["procs"],
            funcs=fx["funcs"], ls=fx["lschemas"], tech=fx["technos"], kms=fx["kms"],
            pattern=load_pattern(fx["writes"]), vars_=fx["vars"], steps=len(s["steps"]),
            tasks=len(s["tasks"]), flags=flags, git=s["git"], odi_date=s["last_date"],
            odi_user=s["last_user"], path=s["file"], releases=s["releases"], criticality=crit,
            component="ODI_SCENARIO"))
    out["cp360_ODI_360"] = cp
    return out


CP360_COLUMNS = [
    "Env", "Layer", "Object_Type", "Object_Name", "Version", "Business_Domain", "Description",
    "Source_Object_Type", "Source_Object_Name", "Load_Plans", "Load_Plan_Step_Path",
    "Autosys_Jobs", "Autosys_Box", "Schedule", "Upstream_Scenarios", "Downstream_Scenarios",
    "Scenarios_In_Load_Plan", "Source_Tables", "Target_Tables", "DB_Links", "DB_Procedures_Called",
    "DB_Functions_Called", "Logical_Schemas", "Technologies", "Knowledge_Modules", "Load_Pattern",
    "Variables", "Step_Count", "Task_Count", "Jira_Refs", "Last_Commit_Date", "Last_Commit_Author",
    "Last_Change_Summary", "ODI_Last_Modified", "ODI_Last_User", "Releases", "Repo_Path",
    "Risk_Flags", "Criticality", "Status_or_Notes", "Project", "Zone", "Component", "Display_Order",
]


def _cp_row(args, order, otype, name, ver, desc, domain="", layer="", src_type="", src_name="",
            load_plans="", lp_path="", jobs=(), box="", sched="", up=(), down=(), src=(), tgt=(),
            dblinks=(), procs=(), funcs=(), ls=(), tech=(), kms=(), pattern="", vars_=(), steps="",
            tasks="", flags=(), git=None, odi_date="", odi_user="", path="", releases=(),
            criticality="", component="", scen_in=()):
    git = git or {}
    return dict(zip(CP360_COLUMNS, [
        args.env, layer, otype, name, ver, domain or "Unclassified", (desc or "")[:1000], src_type, src_name,
        load_plans, lp_path, join(j["job"] for j in jobs), box, sched, join(up), join(down), join(scen_in),
        join(src), join(tgt), join(dblinks), join(procs), join(funcs), join(ls), join(tech), join(kms),
        pattern, join(vars_), steps, tasks, join(git.get("jira", [])), git.get("last_date", ""),
        git.get("last_author", ""), git.get("last_subject", ""), odi_date, odi_user, join(releases), path,
        join(flags), criticality, "Auto-extracted by odi360_extract.py v" + VERSION, args.project, "DATA",
        component, order]))


# ---------------------------------------------------------------- writers
def write_csv(path, rows, cols=None):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    cols = cols or _cols(rows)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({k: _cell(r.get(k, "")) for k in cols})


def _cols(rows):
    cols = []
    for r in rows:
        for k in r:
            if k not in cols:
                cols.append(k)
    return cols or ["empty"]


def _cell(v):
    if isinstance(v, (set, list, tuple)):
        v = join(v)
    elif isinstance(v, dict):
        v = json.dumps(v, default=str)[:MAX_CELL]
    v = "" if v is None else v
    if isinstance(v, str):
        v = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", v)[:MAX_CELL]
    return v


def write_xlsx(path, sheets):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return False
    wb = Workbook()
    wb.remove(wb.active)
    for name, (rows, cols) in sheets.items():
        ws = wb.create_sheet(name[:31])
        cols = cols or _cols(rows)
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
        for r in rows:
            ws.append([_cell(r.get(k, "")) for k in cols])
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        for i, k in enumerate(cols, 1):
            width = max([len(str(k))] + [min(len(str(_cell(r.get(k, "")))), 60) for r in rows[:200]])
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 62)
    wb.save(path)
    return True


def write_knowledge(outdir, args, data, stats, part_limit=350_000):
    """Markdown knowledge base, chunked for Claude Enterprise project upload."""
    kdir = os.path.join(outdir, "claude_ingest")
    os.makedirs(kdir, exist_ok=True)
    sections = []
    head = [f"# ODI 360 knowledge base - {args.project} / {args.env}",
            f"Generated {datetime.datetime.now():%Y-%m-%d %H:%M} by odi360_extract.py v{VERSION}.",
            "Facts are machine-extracted from Bitbucket; heuristic fields are marked (heuristic).", "",
            "## Inventory", *[f"- {k}: {v}" for k, v in stats.items()], ""]
    sections.append("\n".join(head))
    lp_steps = defaultdict(list)
    for st in data["loadplan_steps"]:
        lp_steps[st["load_plan"]].append(st)
    for lp in data["loadplans"]:
        lines = [f"## LOAD PLAN {lp['load_plan']}",
                 f"- Description: {lp['description'] or 'n/a'}",
                 f"- Autosys jobs: {lp['autosys_jobs'] or 'none found'}; schedule: {lp['schedule'] or 'n/a'}",
                 f"- Scenarios ({lp['scenario_count']}): {lp['scenarios']}",
                 f"- Disabled steps: {lp['disabled_steps']}; variables: {lp['variables'] or 'none'}",
                 f"- Last change: {lp['git_last_date'] or 'n/a'}; Jira: {lp['jira'] or 'none'}; file: {lp['file']}",
                 "- Step tree:"]
        lines += [f"  - {s['path']}" + (f" -> runs {s['scenario']} v{s['scenario_version'] or 'latest'}" if s["scenario"] else "")
                  + (" (DISABLED)" if s["enabled"] == "N" else "") for s in lp_steps[lp["load_plan"]]]
        sections.append("\n".join(lines))
    cp = {(r["Object_Name"], r["Version"]): r for r in data["cp360_ODI_360"] if r["Object_Type"] == "SCENARIO"}
    for s in data["scenarios"]:
        c = cp.get((s["scenario"], s["version"]), {})
        sections.append("\n".join([
            f"## SCENARIO {s['scenario']} v{s['version']}",
            f"- Description: {s['description'] or 'n/a'}",
            f"- Generated from: {s['source_object_type']} {s['source_object_name']} (folder {s['source_folder'] or 'n/a'})",
            f"- Orchestration: load plans [{s['load_plans'] or 'none'}]; Autosys [{s['autosys_jobs'] or 'none'}]; scripts [{s['scripts'] or 'none'}]",
            f"- Schedule: {c.get('Schedule') or 'n/a'}",
            f"- Reads: {s['reads'] or 'none resolved'}",
            f"- Writes: {s['writes'] or 'none resolved'}",
            f"- Load pattern: {s['load_pattern']}; KMs: {s['kms'] or 'n/a'}",
            f"- DB procedures: {s['db_procedures'] or 'none'}; functions: {s['db_functions'] or 'none'}; DB links: {s['db_links'] or 'none'}",
            f"- Upstream scenarios (data): {c.get('Upstream_Scenarios') or 'none'}; downstream: {c.get('Downstream_Scenarios') or 'none'}",
            f"- Logical schemas: {s['logical_schemas'] or 'n/a'}; technologies: {s['technologies'] or 'n/a'}",
            f"- Variables: {s['variables'] or 'none'}",
            f"- Steps ({s['step_count']}): {s['steps'][:1500] or 'n/a'}",
            f"- Domain (heuristic): {c.get('Business_Domain', '')}; layer (heuristic): {c.get('Layer', '')}; criticality (heuristic): {c.get('Criticality', '')}",
            f"- Risk flags: {c.get('Risk_Flags') or 'none'}",
            f"- Git: last {s['git_last_date'] or 'n/a'} by {s['git_last_author'] or 'n/a'} - '{s['git_last_subject'] or 'n/a'}'; Jira: {s['jira'] or 'none'}; releases: {s['releases'] or 'none'}",
            f"- ODI last modified: {s['odi_last_date']} by {s['odi_last_user']}; file: {s['file']}"]))
    for t in data["table_360"]:
        sections.append(f"## TABLE {t['table']}\n- Type/schema: {t['db_object_type'] or 'not in DB scripts'} {t['schema']}\n"
                        f"- Written by: {t['written_by_scenarios'] or 'none'} ({t['write_operations']})\n"
                        f"- Read by: {t['read_by_scenarios'] or 'none'}\n- DB code using it: {t['db_objects_using'] or 'none'}\n"
                        f"- Indexes: {t['indexes'] or 'none'}")
    dbc = [d for d in data["db_objects"] if d["object_type"] in ("PROCEDURE", "FUNCTION", "PACKAGE", "PACKAGE BODY", "VIEW", "MATERIALIZED VIEW", "TRIGGER")]
    for d in dbc:
        sections.append(f"## DB {d['object_type']} {d['schema']}.{d['name']}\n- Reads: {d['reads'] or 'none'}\n"
                        f"- Writes: {d['writes'] or 'none'}\n- Calls: {d['calls'] or 'none'}\n"
                        f"- Used by: {d['used_by'] or 'nothing in scope'}\n- File: {d['repo']}/{d['file']} (last {d['last_commit_date']}, Jira {d['jira'] or 'none'})")
    if data["findings"]:
        sections.append("## FINDINGS\n" + "\n".join(f"- [{f['severity']}] {f['code']} {f['object']}: {f['detail']}"
                                                      for f in data["findings"][:3000]))
    parts, cur = [], ""
    for s in sections:
        if len(cur) + len(s) > part_limit and cur:
            parts.append(cur)
            cur = ""
        cur += s + "\n\n"
    parts.append(cur)
    names = []
    for i, p in enumerate(parts, 1):
        fn = f"ODI360_KNOWLEDGE_part{i:02d}.md"
        with open(os.path.join(kdir, fn), "w", encoding="utf-8") as f:
            f.write(p)
        names.append(fn)
    nodes = {}
    for e in data["lineage_edges"]:
        for t, n in ((e["src_type"], e["src"]), (e["dst_type"], e["dst"])):
            nodes.setdefault(f"{t}:{n}", {"id": f"{t}:{n}", "type": t, "name": n})
    with open(os.path.join(kdir, "odi360_graph.json"), "w", encoding="utf-8") as f:
        json.dump({"nodes": list(nodes.values()), "edges": data["lineage_edges"]}, f, indent=1, default=str)
    return names


# ---------------------------------------------------------------- main
def autodetect(root):
    found = {"odi": [], "db": [], "autosys": [], "scripts": []}
    for d in sorted(os.listdir(root)):
        p = os.path.join(root, d)
        if not os.path.isdir(p) or d in SKIP_DIRS:
            continue
        low = d.lower().replace("_", "-").replace(" ", "-")
        subs = {x.lower() for x in os.listdir(p)}
        if "oracle-data-integrator" in low or "odi" == low or \
                {"master repository", "work repository"} & subs or "scenarios and load plans" in subs:
            found["odi"].append(p)
        elif "db-scripts" in low or "database scripts" in subs:
            found["db"].append(p)
        elif "autosys" in low:
            found["autosys"].append(p)
        else:
            found["scripts"].append(p)
    return found


def main(argv=None):
    ap = argparse.ArgumentParser(description="ODI 360 data-point extractor")
    ap.add_argument("--root", help="folder holding the Bitbucket clones (auto-detects repos)")
    ap.add_argument("--odi-repo", action="append", default=[], help="oracle-data-integrator clone")
    ap.add_argument("--db-scripts", action="append", default=[], help="prod-db-scripts clone (repeatable)")
    ap.add_argument("--autosys", action="append", default=[], help="Autosys JIL repo (repeatable)")
    ap.add_argument("--scripts", action="append", default=[], help="other repos to scan (repeatable)")
    ap.add_argument("--out", default="odi360_out")
    ap.add_argument("--env", default="PROD", help="Env column value (master branch = PROD)")
    ap.add_argument("--project", default="PBDW")
    ap.add_argument("--no-git", action="store_true", help="skip git log metadata")
    args = ap.parse_args(argv)

    if args.root:
        det = autodetect(args.root)
        args.odi_repo += det["odi"]
        args.db_scripts += det["db"]
        args.autosys += det["autosys"]
        args.scripts += det["scripts"]
    if not args.odi_repo:
        ap.error("no ODI repo found - pass --odi-repo or --root")
    for k in ("odi_repo", "db_scripts", "autosys", "scripts"):
        setattr(args, k, [os.path.abspath(p) for p in dict.fromkeys(getattr(args, k))])
    log = lambda m: print(f"[odi360] {m}", file=sys.stderr)
    log(f"ODI={args.odi_repo} DB={args.db_scripts} AUTOSYS={args.autosys} SCRIPTS={args.scripts}")

    git = GitIndex(enabled=not args.no_git)
    for p in args.odi_repo + args.db_scripts + args.autosys + args.scripts:
        git.index(p)

    model = ODIModel()
    for p in args.odi_repo:
        load_odi(p, git, model)
    master_rows, work_rows, var_rows = build_odi(model, git, args.odi_repo[0])
    log(f"scenarios={len(model.scenarios)} loadplans={len(model.loadplans)} work_objects={len(work_rows)}")
    db_objs, fks, _ = load_db_scripts(args.db_scripts, git)
    log(f"db objects={len(db_objs)}")
    jobs = load_autosys(args.autosys + args.scripts, git)
    log(f"autosys jobs={len(jobs)}")
    callables = {d["name"] for d in db_objs if d["object_type"] in ("PROCEDURE", "FUNCTION", "PACKAGE")}
    script_rows = scan_scripts(args.scripts, git, {k[0] for k in model.scenarios}, set(model.loadplans), callables)
    data = build_360(args, model, master_rows, work_rows, var_rows, db_objs, fks, jobs, script_rows)

    out = os.path.abspath(args.out)
    dp = os.path.join(out, "datapoints")
    order = ["scenarios", "scenario_tasks", "loadplans", "loadplan_steps", "table_usage", "table_360",
             "scenario_dependencies", "db_objects", "db_foreign_keys", "autosys_jobs", "script_references",
             "odi_topology", "odi_work_objects", "odi_variables", "releases", "lineage_edges", "findings",
             "parse_errors"]
    sheets = {}
    for k in order:
        write_csv(os.path.join(dp, f"{k}.csv"), data[k])
        sheets[k] = (data[k], None)
    stats = {k: len(data[k]) for k in order}
    stats.update({"findings_" + s: sum(1 for f in data["findings"] if f["severity"] == s)
                  for s in ("HIGH", "MEDIUM", "LOW", "INFO")})
    write_csv(os.path.join(out, "cp360", "ODI_360.csv"), data["cp360_ODI_360"], CP360_COLUMNS)
    xl = write_xlsx(os.path.join(out, "ODI360_datapoints.xlsx"), sheets)
    write_xlsx(os.path.join(out, "cp360", "ODI_360.xlsx"),
               {"ODI_360": (data["cp360_ODI_360"], CP360_COLUMNS),
                "ODI_360_LINEAGE": (data["lineage_edges"], None),
                "ODI_360_TABLES": (data["table_360"], None),
                "ODI_360_FINDINGS": (data["findings"], None)})
    parts = write_knowledge(out, args, data, stats)
    manifest = {"generated": datetime.datetime.now().isoformat(timespec="seconds"), "version": VERSION,
                "env": args.env, "project": args.project, "inputs": {k: getattr(args, k) for k in
                                                                     ("odi_repo", "db_scripts", "autosys", "scripts")},
                "stats": stats, "knowledge_files": parts, "xlsx": xl}
    with open(os.path.join(out, "claude_ingest", "odi360_manifest.json"), "w", encoding="utf-8") as f:
        json.dump(manifest, f, indent=2)
    log(json.dumps(stats))
    log(f"done -> {out}" + ("" if xl else " (pip install openpyxl for .xlsx output)"))
    return 0


if __name__ == "__main__":
    sys.exit(main())
